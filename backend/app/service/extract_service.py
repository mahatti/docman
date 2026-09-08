import html
import math
import re
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

CHARS_PER_PAGE = 2500
W_LAST_RENDERED_PAGE_BREAK = qn("w:lastRenderedPageBreak")
W_BR = qn("w:br")
W_TYPE = qn("w:type")
W_T = qn("w:t")
APP_PROPS_NS = {"ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"}

STOPWORDS = {"the", "and", "untuk", "yang", "dari", "dengan", "pada", "atau"}
GENERIC_MODULES = {"umum"}

SEGMENT_HEADING_RE = re.compile(
    r"(?im)^\s*(?:\d+(?:\.\d+)*\.?\s+)?segments?\s+([A-Z]{2,}(?:[-_][A-Za-z0-9_]+)+)\s*$"
)
SEGMENT_INLINE_RE = re.compile(
    r"(?i)\bsegments?\s+([A-Z]{2,}(?:[-_][A-Za-z0-9_]+)+)"
)
PROC_AFTER_SEGMENT_RE = re.compile(
    r"(?i)\bsegments?\s+[A-Z0-9_\-]+\s+[–—-]\s+([A-Za-z][A-Za-z0-9_]*)"
)
SQL_PROC_RE = re.compile(
    r"\b((?:usp[i]?|sp)_[A-Za-z][A-Za-z0-9_]*)\b",
    re.IGNORECASE,
)
SPEC_TABLE_RE = re.compile(r"(?i)list of specification table for\s+([A-Za-z][A-Za-z0-9_\.]*)")
APPENDIX_SOURCE_RE = re.compile(r"(?i)^3\.b\.1\s+table source\s*$")
APPENDIX_PARAM_RE = re.compile(r"(?i)^3\.b\.2\s+table parameter")
APPENDIX_TARGET_RE = re.compile(r"(?i)^3\.b\.3\s+table target\s*$")
FLOW_SOURCE_RE = re.compile(r"(?i)^(?:🟦\s*)?table source\s*$")
FLOW_TARGET_RE = re.compile(r"(?i)^(?:🟥\s*)?target table\s*$")
FLOW_PROCESS_RE = re.compile(r"(?i)^(?:🟧\s*)?proses transformasi")
FLOW_PROC_RE = re.compile(
    r"(?i)data (?:models|flow).+segment\s+\S+\s+[–—-]\s+([A-Za-z][A-Za-z0-9_]*)\s*$"
)
FLOW_TABLE_LINE_RE = re.compile(
    r"^([A-Za-z][A-Za-z0-9_\.]{2,})(?:\s*\([^)]*\))?\s*:\s+(.+)$"
)
SKIP_TABLE_NAMES = {
    "fungsi",
    "isi",
    "tujuan",
    "updated",
    "existing",
    "field",
    "datatype",
    "description",
    "deskripsi",
}
DOC_VERSION_RE = re.compile(r"(?i)document\s+version\s+([0-9]+(?:\.[0-9]+)*)")
VERSION_TOKEN_RE = re.compile(r"^[0-9]+(?:\.[0-9]+)+$")
FILENAME_TSD_RE = re.compile(
    r"(?i)(?:^|[\s_])([A-Z]{2,})_TSD[_-]([A-Za-z0-9_]+)"
)
TITLE_AFTER_FOR_RE = re.compile(
    r"(?is)technical specification document(?:\s+for)?\s+([^\n]+)"
)


def tags_from_filename(name: str) -> list[str]:
    stem = Path(name).stem
    parts = re.split(r"[_\-\s]+", stem)
    tags = []
    seen = set()
    for part in parts:
        tag = re.sub(r"[^a-zA-Z0-9]+", "", part).lower()
        if len(tag) < 2 or tag in seen or tag in STOPWORDS:
            continue
        seen.add(tag)
        tags.append(tag)
        if len(tags) >= 6:
            break
    return tags or ["baru"]


def count_procedures(text: str) -> int:
    names = extract_procedure_names(text or "")
    if names:
        return len(names)
    numbered = re.findall(
        r"(?m)^\s*(?:\d{1,2}[\.\)]\s+\S+|(?:langkah|prosedur|procedure)\s+\d+)",
        text or "",
        flags=re.I,
    )
    labeled = re.findall(r"(?i)\b(?:prosedur|sop)\b", text or "")
    return max(len(numbered), min(len(labeled), 40))


def iter_docx_blocks(document: DocxDocument):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _is_page_break(element) -> bool:
    if element.tag == W_LAST_RENDERED_PAGE_BREAK:
        return True
    return element.tag == W_BR and element.get(W_TYPE) == "page"


def _paragraph_line(paragraph: Paragraph) -> str:
    return (paragraph.text or "").strip()


def _table_lines(table: Table) -> list[str]:
    lines = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells if cell.text.strip()]
        if cells:
            lines.append(" || ".join(cells))
    return lines


def read_docx_text(file_path: str) -> str:
    document = DocxDocument(file_path)
    lines = []
    for block in iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            text = _paragraph_line(block)
            if text:
                lines.append(text)
            continue
        lines.extend(_table_lines(block))
    return "\n".join(lines)


def _docx_app_page_count(file_path: str) -> int | None:
    try:
        with ZipFile(file_path) as archive:
            xml = archive.read("docProps/app.xml")
        root = ET.fromstring(xml)
        node = root.find("ep:Pages", APP_PROPS_NS)
        if node is None or not node.text:
            return None
        count = int(node.text)
        return count if count > 0 else None
    except (OSError, ET.ParseError, ValueError, KeyError):
        return None


def _split_text_into_pages(text: str, total_pages: int) -> list[tuple[int, str]]:
    content = text or ""
    if total_pages <= 1:
        return [(1, content)] if content.strip() else []
    if "\f" in content:
        parts = [part.strip() for part in content.split("\f")]
        pages = [(index, part) for index, part in enumerate(parts, start=1) if part]
        return pages or [(1, content)]
    length = len(content)
    pages: list[tuple[int, str]] = []
    for index in range(total_pages):
        start = (index * length) // total_pages
        end = ((index + 1) * length) // total_pages
        chunk = content[start:end].strip()
        if chunk:
            pages.append((index + 1, chunk))
    return pages


def _paragraph_page_segments(paragraph: Paragraph) -> list[tuple[bool, str]]:
    segments: list[tuple[bool, str]] = []
    buffer: list[str] = []
    starts_new_page = False

    def flush():
        nonlocal starts_new_page
        text = "".join(buffer)
        buffer.clear()
        stripped = text.strip()
        if stripped or starts_new_page:
            segments.append((starts_new_page, stripped))
        starts_new_page = False

    for run in paragraph.runs:
        for child in run._element.iter():
            if child is run._element:
                continue
            if _is_page_break(child):
                flush()
                starts_new_page = True
            elif child.tag == W_T and child.text:
                buffer.append(child.text)
    flush()
    if not segments:
        text = _paragraph_line(paragraph)
        if text:
            segments.append((False, text))
    return segments


def _element_has_page_break(element) -> bool:
    return any(_is_page_break(child) for child in element.iter())


def read_docx_pages(file_path: str) -> tuple[int, list[tuple[int, str]]]:
    document = DocxDocument(file_path)
    buckets: list[list[str]] = [[]]

    def add_line(text: str):
        if text:
            buckets[-1].append(text)

    def start_new_page():
        buckets.append([])

    for block in iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            for starts_new, text in _paragraph_page_segments(block):
                if starts_new:
                    start_new_page()
                add_line(text)
            continue
        for row in block.rows:
            if _element_has_page_break(row._tr):
                start_new_page()
            cells = [" ".join(cell.text.split()) for cell in row.cells if cell.text.strip()]
            if cells:
                add_line(" || ".join(cells))

    raw_pages = ["\n".join(lines).strip() for lines in buckets]
    while len(raw_pages) > 1 and not raw_pages[-1]:
        raw_pages.pop()
    break_count = max(len(raw_pages), 1)
    declared = _docx_app_page_count(file_path)
    total = max(break_count, declared or 0, 1)
    nonempty = [(index, text) for index, text in enumerate(raw_pages, start=1) if text]
    if len(nonempty) <= 1 and total > 1:
        full_text = nonempty[0][1] if nonempty else read_docx_text(file_path)
        nonempty = _split_text_into_pages(full_text, total)
    return total, nonempty


def read_pdf_pages(file_path: str) -> tuple[int, list[tuple[int, str]]]:
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    total = len(reader.pages)
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append((index, text))
    return max(total, 1), pages


def read_txt_pages(file_path: str) -> tuple[int, list[tuple[int, str]]]:
    text = Path(file_path).read_text(encoding="utf-8", errors="ignore")
    if "\f" in text:
        parts = [part.strip() for part in text.split("\f")]
        pages = [(index, part) for index, part in enumerate(parts, start=1) if part]
        return max(len(parts), 1), pages
    total = max(1, math.ceil(len(text.strip()) / CHARS_PER_PAGE)) if text.strip() else 1
    return total, _split_text_into_pages(text, total)


def load_paged_text(file_path: str, file_type: str | None = None) -> tuple[int, list[tuple[int, str]]]:
    ext = (file_type or Path(file_path).suffix.lstrip(".")).lower()
    if ext == "pdf":
        return read_pdf_pages(file_path)
    if ext == "docx":
        return read_docx_pages(file_path)
    if ext == "txt":
        return read_txt_pages(file_path)
    if ext == "doc":
        raise ValueError("Format .doc lama tidak didukung. Unggah ulang sebagai .docx.")
    raise ValueError(f"Tipe file tidak didukung: {ext}")


def page_for_content(content: str, pages: list[tuple[int, str]]) -> int:
    if not pages:
        return 1
    needle = re.sub(r"\s+", " ", content or "").strip()
    if not needle:
        return pages[0][0]
    normalized = [(page, re.sub(r"\s+", " ", text or "")) for page, text in pages]
    for size in (180, 80, 40):
        snippet = needle[:size]
        matches = [page for page, text in normalized if snippet and snippet in text]
        if matches:
            return matches[0]
    best_page, best_score = pages[0][0], -1
    sample = needle[:24]
    for page, text in normalized:
        score = text.find(sample) if sample else -1
        if score >= 0 and (best_score < 0 or score < best_score):
            best_page, best_score = page, score
    return best_page


def read_file_text(file_path: str, file_type: str | None = None) -> str:
    ext = (file_type or Path(file_path).suffix.lstrip(".")).lower()
    if ext == "docx":
        return read_docx_text(file_path)
    if ext == "txt":
        return Path(file_path).read_text(encoding="utf-8", errors="ignore")
    if ext == "pdf":
        _, pages = read_pdf_pages(file_path)
        return "\n".join(text for _, text in pages)
    if ext == "doc":
        raise ValueError("Format .doc lama tidak didukung. Unggah ulang sebagai .docx.")
    raise ValueError(f"Tipe file tidak didukung: {ext}")


def _run_html(run) -> str:
    text = html.escape(run.text or "")
    if not text:
        return ""
    if run.bold:
        text = f"<strong>{text}</strong>"
    if run.italic:
        text = f"<em>{text}</em>"
    if run.underline:
        text = f"<u>{text}</u>"
    return text


def _paragraph_html(paragraph: Paragraph) -> str:
    inner = "".join(_run_html(run) for run in paragraph.runs).strip()
    if not inner:
        inner = html.escape(_paragraph_line(paragraph))
    if not inner:
        return ""
    style = ((paragraph.style.name if paragraph.style else "") or "").lower()
    if style.startswith("heading 1") or style == "title":
        return f"<h1>{inner}</h1>"
    if style.startswith("heading 2") or style == "subtitle":
        return f"<h2>{inner}</h2>"
    if style.startswith("heading 3"):
        return f"<h3>{inner}</h3>"
    if style.startswith("heading 4"):
        return f"<h4>{inner}</h4>"
    if "list number" in style or style.startswith("list number"):
        return f"<li data-list='ol'>{inner}</li>"
    if "list" in style:
        return f"<li data-list='ul'>{inner}</li>"
    return f"<p>{inner}</p>"


def _table_html(table: Table) -> str:
    rows_html = []
    for index, row in enumerate(table.rows):
        cell_tag = "th" if index == 0 else "td"
        cells = "".join(
            f"<{cell_tag}>{html.escape(' '.join(cell.text.split()))}</{cell_tag}>"
            for cell in row.cells
        )
        rows_html.append(f"<tr>{cells}</tr>")
    if not rows_html:
        return ""
    return f"<table><tbody>{''.join(rows_html)}</tbody></table>"


def _flush_list(items: list[str]) -> str:
    if not items:
        return ""
    numbered = any('data-list="ol"' in item or "data-list='ol'" in item for item in items)
    tag = "ol" if numbered else "ul"
    cleaned = [re.sub(r"\sdata-list=['\"][ou]l['\"]", "", item) for item in items]
    items.clear()
    return f"<{tag}>{''.join(cleaned)}</{tag}>"


def docx_to_html(file_path: str) -> str:
    document = DocxDocument(file_path)
    parts: list[str] = []
    list_items: list[str] = []
    for block in iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            markup = _paragraph_html(block)
            if markup.startswith("<li"):
                list_items.append(markup)
                continue
            if list_items:
                parts.append(_flush_list(list_items))
            if markup:
                parts.append(markup)
            continue
        if list_items:
            parts.append(_flush_list(list_items))
        table_markup = _table_html(block)
        if table_markup:
            parts.append(table_markup)
    if list_items:
        parts.append(_flush_list(list_items))
    return "".join(parts) or "<p>Dokumen kosong.</p>"


def document_to_preview_html(file_path: str, file_type: str | None = None) -> str:
    ext = (file_type or Path(file_path).suffix.lstrip(".")).lower()
    if ext == "docx":
        return docx_to_html(file_path)
    text = read_file_text(file_path, ext)
    paragraphs = [f"<p>{html.escape(line)}</p>" for line in text.splitlines() if line.strip()]
    return "".join(paragraphs) or "<p>Dokumen kosong.</p>"


def _normalize_segment(value: str | None) -> str | None:
    if not value:
        return None
    cleaned = re.sub(r"\s+", "", value).strip(" .-_")
    cleaned = re.sub(r"\d+$", "", cleaned)
    return cleaned or None


def extract_document_code(text: str, filename: str = "") -> str | None:
    for match in SEGMENT_HEADING_RE.finditer(text or ""):
        code = _normalize_segment(match.group(1))
        if code:
            return code
    counts: dict[str, int] = {}
    for match in SEGMENT_INLINE_RE.finditer(text or ""):
        code = _normalize_segment(match.group(1))
        if not code:
            continue
        counts[code] = counts.get(code, 0) + 1
    if counts:
        return sorted(counts.items(), key=lambda item: (-item[1], -len(item[0])))[0][0]

    file_match = FILENAME_TSD_RE.search(Path(filename).stem if filename else "")
    if file_match:
        module, rest = file_match.groups()
        return f"{module}-{rest}"
    return None


def extract_module_name(document_code: str | None, text: str = "", filename: str = "") -> str | None:
    if document_code:
        return document_code.split("-", 1)[0].split("_", 1)[0].upper()
    file_match = FILENAME_TSD_RE.search(Path(filename).stem if filename else "")
    if file_match:
        return file_match.group(1).upper()
    heading = SEGMENT_HEADING_RE.search(text or "")
    if heading:
        return heading.group(1).split("-", 1)[0].upper()
    return None


def extract_version(text: str) -> str | None:
    versions = []
    in_change_control = False
    for raw_line in (text or "").splitlines():
        line = raw_line.strip()
        cells = [part.strip() for part in re.split(r"\s*\|\|\s*|\t+", line) if part.strip()]
        if cells and cells[0].lower() == "version":
            in_change_control = True
            continue
        if in_change_control:
            if not cells or not VERSION_TOKEN_RE.match(cells[0]):
                if versions:
                    break
                continue
            versions.append(cells[0])
    if versions:
        return versions[-1]
    header = DOC_VERSION_RE.search(text or "")
    if header:
        return header.group(1)
    return None


def extract_title(text: str, document_code: str | None, filename: str = "") -> str:
    if document_code:
        return f"TSD {document_code}"
    match = TITLE_AFTER_FOR_RE.search(text or "")
    if match:
        title = match.group(1).strip()
        title = re.split(r"[\n\r]", title, maxsplit=1)[0].strip()
        if title:
            return title
    if filename:
        return Path(filename).stem.replace("_", " ")
    return "TSD"


def _is_heading_like(line: str) -> bool:
    stripped = line.strip()
    if not stripped or len(stripped) > 180:
        return False
    if stripped.startswith(("http://", "https://")):
        return False
    return bool(
        re.match(r"^\d+(?:\.\d+)*", stripped)
        or re.match(r"(?i)^(data models|data flow|start procedure|list of procedures)", stripped)
        or "segment" in stripped.lower()
    )


def _canonical_procedure_names(names: list[str]) -> list[str]:
    dropped = set()
    for current in names:
        for other in names:
            if current.lower() == other.lower():
                continue
            if not current.lower().startswith(other.lower()):
                continue
            suffix = current[len(other) :]
            if not suffix.isdigit():
                continue
            if len(suffix) >= 2:
                dropped.add(current.lower())
            else:
                dropped.add(other.lower())
    return [name for name in names if name.lower() not in dropped]


def extract_procedure_names(text: str) -> list[str]:
    names: list[str] = []
    seen = set()

    def add(name: str | None):
        cleaned = (name or "").strip().strip(".-")
        if len(cleaned) < 4 or cleaned.lower() in seen:
            return
        if not re.search(r"[A-Za-z]", cleaned):
            return
        seen.add(cleaned.lower())
        names.append(cleaned)

    for match in PROC_AFTER_SEGMENT_RE.finditer(text or ""):
        add(match.group(1))

    if names:
        return _canonical_procedure_names(names)

    for line in (text or "").splitlines():
        if not _is_heading_like(line):
            continue
        for match in SQL_PROC_RE.finditer(line):
            add(match.group(1))
    return _canonical_procedure_names(names)


def _looks_like_table(name: str) -> bool:
    if "_" in name or "." in name:
        return True
    if re.match(r"(?i)^(tbl|rwa|prc|usp|sp_)", name):
        return True
    return bool(name.isupper() and len(name) >= 4)


def _clean_table_name(name: str | None) -> str | None:
    cleaned = (name or "").strip().strip(".:")
    cleaned = re.sub(r"\s*\([^)]*\)\s*$", "", cleaned).strip()
    if len(cleaned) < 3 or cleaned.lower() in SKIP_TABLE_NAMES:
        return None
    if " " in cleaned:
        return None
    if not re.match(r"^[A-Za-z][A-Za-z0-9_\.]*$", cleaned):
        return None
    if not _looks_like_table(cleaned):
        return None
    return cleaned


def _add_catalog_table(
    bucket: list[dict],
    name: str | None,
    description: str = "",
    procedure: str | None = None,
    section: str = "",
):
    cleaned = _clean_table_name(name)
    if not cleaned:
        return
    desc = re.sub(r"\s+", " ", description or "").strip()
    key = cleaned.lower()
    for item in bucket:
        if item["name"].lower() != key:
            continue
        if desc and len(desc) > len(item.get("description") or ""):
            item["description"] = desc
        if procedure and not item.get("procedure"):
            item["procedure"] = procedure
        return
    bucket.append(
        {
            "name": cleaned,
            "description": desc[:280],
            "procedure": procedure,
            "section": section,
        }
    )


def _parse_appendix_tables(lines: list[str], start: int, stop_at) -> tuple[list[dict], int]:
    items: list[dict] = []
    index = start
    while index < len(lines):
        line = lines[index].strip()
        if stop_at(line):
            break
        match = SPEC_TABLE_RE.search(line)
        if match:
            _add_catalog_table(items, match.group(1), section="appendix")
        index += 1
    return items, index


def _parse_flow_tables(lines: list[str], start: int) -> tuple[list[dict], int]:
    items: list[dict] = []
    index = start
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if (
            FLOW_SOURCE_RE.match(line)
            or FLOW_TARGET_RE.match(line)
            or FLOW_PROCESS_RE.match(line)
            or FLOW_PROC_RE.search(line)
            or APPENDIX_SOURCE_RE.match(line)
            or APPENDIX_TARGET_RE.match(line)
            or APPENDIX_PARAM_RE.match(line)
            or re.match(r"(?i)^3\.\s*appendix", line)
        ):
            break
        match = FLOW_TABLE_LINE_RE.match(line)
        if match:
            _add_catalog_table(items, match.group(1), match.group(2), section="data_flow")
        index += 1
    return items, index


def extract_table_catalog(text: str) -> dict:
    lines = (text or "").splitlines()
    source: list[dict] = []
    target: list[dict] = []
    current_procedure = None
    last_source_for_proc: list[dict] = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        proc_match = FLOW_PROC_RE.search(line)
        if proc_match:
            current_procedure = proc_match.group(1)
            last_source_for_proc = []
            index += 1
            continue
        if APPENDIX_SOURCE_RE.match(line):
            items, index = _parse_appendix_tables(
                lines,
                index + 1,
                lambda value: bool(APPENDIX_PARAM_RE.match(value) or APPENDIX_TARGET_RE.match(value)),
            )
            for item in items:
                _add_catalog_table(source, item["name"], item.get("description") or "", section="appendix")
            continue
        if APPENDIX_TARGET_RE.match(line):
            items, index = _parse_appendix_tables(
                lines,
                index + 1,
                lambda value: bool(re.match(r"(?i)^3\.[a-z]", value) and not value.lower().startswith("3.b.3")),
            )
            for item in items:
                _add_catalog_table(target, item["name"], item.get("description") or "", section="appendix")
            continue
        if FLOW_SOURCE_RE.match(line):
            items, index = _parse_flow_tables(lines, index + 1)
            last_source_for_proc = items
            for item in items:
                _add_catalog_table(
                    source,
                    item["name"],
                    item.get("description") or "",
                    current_procedure,
                    "data_flow",
                )
            continue
        if FLOW_TARGET_RE.match(line):
            items, index = _parse_flow_tables(lines, index + 1)
            if not items and last_source_for_proc:
                items = [
                    {
                        **item,
                        "description": item.get("description") or "Tabel yang sama di-update sebagai target.",
                    }
                    for item in last_source_for_proc
                ]
            for item in items:
                _add_catalog_table(
                    target,
                    item["name"],
                    item.get("description") or "",
                    current_procedure,
                    "data_flow",
                )
            continue
        index += 1
    return {"source": source, "target": target}


def table_question_role(question: str) -> str | None:
    text = (question or "").lower()
    if not re.search(r"\b(tabel|table)\b", text):
        return None
    wants_source = bool(re.search(r"\b(sumber|source)\b", text))
    wants_target = bool(re.search(r"\b(target|tujuan)\b", text))
    if wants_source and not wants_target:
        return "source"
    if wants_target and not wants_source:
        return "target"
    return "both"


def format_table_catalog(catalog: dict | None, question: str = "", document_name: str = "") -> str:
    if not catalog:
        return ""
    role = table_question_role(question) or "both"
    q = (question or "").lower()
    matched_proc = None
    for item in (catalog.get("source") or []) + (catalog.get("target") or []):
        proc = (item.get("procedure") or "").strip()
        if proc and proc.lower() in q:
            matched_proc = proc
            break

    def select(kind: str) -> list[dict]:
        rows = catalog.get(kind) or []
        if matched_proc:
            filtered = [row for row in rows if (row.get("procedure") or "").lower() == matched_proc.lower()]
            if filtered:
                return filtered
        return rows

    sections = []
    if role in {"source", "both"}:
        rows = select("source")
        if rows:
            heading = f"Tabel sumber (Table Source) di {document_name or 'dokumen'}"
            if matched_proc:
                heading += f" untuk prosedur {matched_proc}"
            lines = [heading + ":"]
            for row in rows:
                extra = f" — {row['description']}" if row.get("description") else ""
                lines.append(f"- **{row['name']}**{extra}")
            sections.append("\n".join(lines))
    if role in {"target", "both"}:
        rows = select("target")
        if rows:
            heading = f"Tabel target (Target Table) di {document_name or 'dokumen'}"
            if matched_proc:
                heading += f" untuk prosedur {matched_proc}"
            lines = [heading + ":"]
            for row in rows:
                extra = f" — {row['description']}" if row.get("description") else ""
                lines.append(f"- **{row['name']}**{extra}")
            sections.append("\n".join(lines))
    return "\n\n".join(sections)


def annotate_table_labels(text: str) -> str:
    updated = re.sub(r"(?i)🟦\s*table source", "TABEL SUMBER (Table Source)", text or "")
    updated = re.sub(r"(?i)🟥\s*target table", "TABEL TARGET (Target Table)", updated)
    updated = re.sub(r"(?im)^3\.b\.1\s+table source", "3.b.1 TABEL SUMBER (Table Source)", updated)
    updated = re.sub(r"(?im)^3\.b\.3\s+table target", "3.b.3 TABEL TARGET (Target Table)", updated)
    return updated


def extract_tsd_metadata(text: str, filename: str = "") -> dict:
    document_code = extract_document_code(text, filename)
    module_name = extract_module_name(document_code, text, filename)
    version = extract_version(text)
    procedures = extract_procedure_names(text)
    title = extract_title(text, document_code, filename)
    tags = []
    if module_name:
        tags.append(module_name.lower())
    if document_code:
        tags.append(document_code.lower())
    tags.extend(item.lower() for item in procedures[:4])
    tags.extend(tags_from_filename(filename))
    unique_tags = []
    seen = set()
    for tag in tags:
        if tag in seen or tag in GENERIC_MODULES:
            continue
        seen.add(tag)
        unique_tags.append(tag)
        if len(unique_tags) >= 8:
            break
    return {
        "title": title,
        "document_code": document_code,
        "module_name": module_name,
        "version": version,
        "procedures": procedures,
        "tables": extract_table_catalog(text),
        "tags": unique_tags or tags_from_filename(filename),
    }


def extract_tsd_from_file(file_path: str, file_type: str | None = None, filename: str = "") -> dict:
    text = read_file_text(file_path, file_type)
    return extract_tsd_metadata(text, filename or Path(file_path).name)
